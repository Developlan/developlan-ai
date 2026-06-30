from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook

from app.models import OpportunityDocument, OpportunityPackage

try:
    import xlrd
except ImportError:
    xlrd = None


UNSUPPORTED_FOR_EXTRACTION = "unsupported for extraction"


def extract_documents(
    package: OpportunityPackage,
) -> OpportunityPackage:

    for document in package.documents:

        if not document.downloaded:
            document.extraction_status = "not downloaded"
            continue

        if not document.local_path:
            document.extracted_text = None
            document.extraction_status = "extraction failed"
            document.extraction_note = "Downloaded document has no local path."
            continue

        try:
            path = Path(document.local_path)

            suffix = path.suffix.lower()

            if suffix == ".pdf":

                document.extracted_text = extract_pdf(path)
                document.extraction_status = "extracted"
                document.extraction_note = None

            elif suffix == ".docx":

                document.extracted_text = extract_docx(path)
                document.extraction_status = "extracted"
                document.extraction_note = None

            elif suffix == ".xlsx":

                document.extracted_text = extract_xlsx(path)
                document.extraction_status = "extracted"
                document.extraction_note = None

            elif suffix == ".xls":

                if xlrd is None:
                    mark_unsupported(
                        document,
                        "Legacy .xls extraction requires xlrd.",
                    )
                else:
                    document.extracted_text = extract_xls(path)
                    document.extraction_status = "extracted"
                    document.extraction_note = None

            elif suffix == ".doc":

                mark_unsupported(
                    document,
                    "Legacy .doc extraction requires external conversion tooling.",
                )

            elif suffix in [".txt", ".html", ".htm"]:

                document.extracted_text = path.read_text(
                    errors="ignore"
                )
                document.extraction_status = "extracted"
                document.extraction_note = None

            else:

                mark_unsupported(
                    document,
                    f"{suffix or 'unknown file type'} is not supported for extraction.",
                )

        except Exception as e:

            document.extracted_text = None
            document.extraction_status = "extraction failed"
            document.extraction_note = str(e)

    return package


def mark_unsupported(
    document: OpportunityDocument,
    note: str,
) -> None:

    document.extracted_text = None
    document.extraction_status = UNSUPPORTED_FOR_EXTRACTION
    document.extraction_note = note


def extract_pdf(path: Path) -> str:

    doc = fitz.open(path)

    pages = []

    for page in doc:

        pages.append(page.get_text())

    doc.close()

    return "\n".join(pages)


def extract_docx(path: Path) -> str:

    doc = Document(path)

    text = []

    for paragraph in doc.paragraphs:

        if paragraph.text.strip():

            text.append(paragraph.text)

    for table in doc.tables:

        for row in table.rows:

            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            text.append(" | ".join(cells))

    return "\n".join(text)


def extract_xlsx(path: Path) -> str:

    workbook = load_workbook(
        path,
        data_only=True,
    )

    output = []

    for sheet in workbook.worksheets:

        output.append(f"Worksheet: {sheet.title}")

        for row in sheet.iter_rows(values_only=True):

            values = [
                str(v)
                for v in row
                if v is not None
            ]

            if values:

                output.append(" | ".join(values))

    return "\n".join(output)


def extract_xls(path: Path) -> str:

    workbook = xlrd.open_workbook(path)

    output = []

    for sheet in workbook.sheets():

        output.append(f"Worksheet: {sheet.name}")

        for row_index in range(sheet.nrows):

            values = [
                str(cell.value)
                for cell in sheet.row(row_index)
                if cell.value not in [None, ""]
            ]

            if values:

                output.append(" | ".join(values))

    return "\n".join(output)
